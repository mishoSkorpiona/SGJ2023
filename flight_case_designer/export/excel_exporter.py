"""
export/excel_exporter.py — Export BOM and cut list as an Excel workbook.

Requires ``openpyxl``.  Install via: pip install openpyxl
"""

from __future__ import annotations

from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from ..core.case_manager import CaseData


def export(filepath: str, data: "CaseData") -> None:
    """Export BOM data to an Excel (.xlsx) file."""
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        raise ImportError(
            "openpyxl is required for Excel export. "
            "Install it via: pip install openpyxl"
        )

    wb = openpyxl.Workbook()

    # --- Panels sheet ---
    ws = wb.active
    ws.title = "Panels"
    headers = ["Name", "Width (mm)", "Height (mm)", "Thickness (mm)", "Material", "Qty"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for panel in data.panels:
        ws.append([
            panel["name"], panel["width_mm"], panel["height_mm"],
            panel["thickness_mm"], panel["material"], panel["qty"],
        ])

    # --- Hardware sheet ---
    ws2 = wb.create_sheet("Hardware")
    hw_headers = ["Name", "Count", "Part Number", "Manufacturer"]
    ws2.append(hw_headers)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for hw in data.hardware_items:
        ws2.append([hw["name"], hw["count"], hw["part_number"], hw["manufacturer"]])

    # --- Profiles sheet ---
    ws3 = wb.create_sheet("Profiles")
    pr_headers = ["Name", "Profile Type", "Length (mm)", "Qty"]
    ws3.append(pr_headers)
    for cell in ws3[1]:
        cell.font = Font(bold=True)
    for cut in data.profile_cuts:
        ws3.append([cut["name"], cut["profile_type"], cut["length_mm"], cut["qty"]])

    # --- Summary sheet ---
    ws4 = wb.create_sheet("Summary")
    ws4.append(["Flight Case BOM Summary"])
    ws4["A1"].font = Font(bold=True, size=14)
    ws4.append(["Project", data.project_name])
    ws4.append(["Case Type", data.case_type])
    ws4.append(["Ext. W × D × H (mm)",
                f"{data.ext_width:.1f} × {data.ext_depth:.1f} × {data.ext_height:.1f}"])
    ws4.append(["Total Weight (kg)", data.total_weight])
    ws4.append(["Est. Cost", data.total_cost])

    wb.save(filepath)
    print(f"Excel BOM exported to {filepath}")
